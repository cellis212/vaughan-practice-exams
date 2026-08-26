"""Parameterize the risk-pooling math block (q18-q21) in the central
"Test Bank Topic 3.xlsx" (banks/intro-to-rmi), leaving the other 46 questions byte-identical.

The block's math, mirroring the original answers:
  - Each person faces loss L (or LC for Person C) with probability p, else $0.
  - Actuarially fair premium = p*L.
  - "Amount of risk" is the coefficient of variation: CV = sqrt((1-p)/p)
    for one insured (original: p=.5 -> 1), CV/sqrt(2) for a two-person pool
    of iid risks (original: "about 0.7"). The dollar SD and the raw loss
    serve as distractors, exactly as in the original numbers.
  - q20's pooled premium is the average of A's and C's fair premiums;
    require() keeps every printed premium a whole-dollar amount.
  - q21 is conceptual (wealth-transfer direction) and stays static; it
    remains in the block so it keeps sharing the b_info scenario.

Run from the repo root:  python3 tools/build_param_topic3_block.py
"""

import os

import openpyxl

BANK = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "banks", "intro-to-rmi", "Test Bank Topic 3.xlsx",
)

PARAMS = (
    "p = pick(0.2, 0.25, 0.4, 0.5)\n"
    "L = pick(4, 5, 6, 10, 20)\n"
    "m = pick(2, 3)\n"
    "LC = m * L\n"
    "require(p*L == round(p*L))\n"          # whole-dollar fair premium
    "require(p*L >= 2)\n"                   # keeps CV != SD (pL=1 collides)
    "cv1 = sqrt((1-p)/p)\n"
    "cv2 = cv1 / sqrt(2)\n"
    "prem_pool = (p*L + p*LC) / 2\n"
    "require(prem_pool == round(prem_pool))"  # whole-dollar pooled premium
)

B_INFO = (
    "\t\\textbf{Questions \\thefirst --\\thelast \\ are based on the "
    "following information:}\n\t\n\t\\bi\n"
    "\t\\item Person A faces a [[pct(1-p,0)]] chance of having a loss of "
    "\\$0 and a [[pct(p,0)]] chance of having a loss of \\$[[num(L)]].  \n"
    "\t\\item Person B faces a [[pct(1-p,0)]] chance of having a loss of "
    "\\$0 and a [[pct(p,0)]] chance of having a loss of \\$[[num(L)]]. \n"
    "\t\\item Person C faces a [[pct(1-p,0)]] chance of having a loss of "
    "\\$0 and a [[pct(p,0)]] chance of having a loss of \\$[[num(LC)]].\n"
    "\t\\ei"
)

# q_id -> field replacements. correct/extras are unchanged from the
# original bank (q18 'ac', q19 'c', q20 'b' — a FALSE-hunt — q21 'cd').
REPLACEMENTS = {
    18: {
        "params": PARAMS,
        "b_info": B_INFO,
        "a": "The actuarially fair premium for Person A is "
             "\\$[[num(p*L)]].",                       # correct
        "b": "The actuarially fair premium for Person A is "
             "\\$[[num(L)]].",                         # forgot the probability
        "c": "The amount of risk the insurer faces with Person A being "
             "their only customer is [[num(cv1,2)]].",  # correct: the CV
        "d": "The amount of risk the insurer faces with Person A being "
             "their only customer is [[num(L*sqrt(p*(1-p)),2)]].",
             # the dollar SD, not the CV
    },
    19: {
        "b": "The amount of risk faced by the insurer is "
             "\\$[[num(p*L)]].",                       # the mean, in dollars
        "c": "The amount of risk faced by the insurer is about "
             "[[num(cv2,2)]].",                        # correct: CV/sqrt(2)
        "d": "The amount of risk faced by the insurer is "
             "\\$[[num(L)]].",                         # the loss itself
    },
    20: {
        "Text": "Suppose that Person C also purchases insurance and joins "
                "Person A in the insurer's risk pool.  Note that the risk "
                "pool now contains \\textbf{Person A and Person C only}.  "
                "The insurer charges both person A and person C a premium "
                "of \\$[[num(prem_pool)]]. Which of the following "
                "statements is FALSE?",
    },
}


def main():
    wb = openpyxl.load_workbook(BANK)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    if "params" not in header:
        ws.cell(row=1, column=len(header) + 1, value="params")
        header.append("params")
    col = {name: i + 1 for i, name in enumerate(header)}

    touched = set()
    for row in ws.iter_rows(min_row=2):
        q_id = row[col["q_id"] - 1].value
        if q_id in REPLACEMENTS:
            for field, value in REPLACEMENTS[q_id].items():
                ws.cell(row=row[0].row, column=col[field], value=value)
            touched.add(q_id)

    missing = set(REPLACEMENTS) - touched
    if missing:
        raise SystemExit("q_ids not found in bank: " + str(sorted(missing)))
    wb.save(BANK)
    print("Parameterized q_ids", sorted(touched), "in", BANK)


if __name__ == "__main__":
    main()
