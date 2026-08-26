"""Python translation of the app.R Shiny server logic shared by the RMI
Intro and Life/Health practice exam generators.

Mirrors, step for step:
  1. Bank loading: purrr::map_df(banks_to_load, read_xlsx + mutate(Topic=...))
  2. filtered_bank(): subset(full_bank, Topic %in% input$sel_topics)
  3. The generate handler: optional set.seed, the not-enough-questions
     check, qind <- sample(1:nrow(cur_bank), nrow(cur_bank)), and the
     create_exam() call (practice mode: 1 version, title "Practice Exam:",
     frontPage = TRUE).

Runs both under Pyodide in the browser and under plain CPython for tests.
"""

import os

import openpyxl

import exam_common
from exam_common import r_char
import engine_intro
import engine_lh
import param_engine

ENGINES = {
    "intro": engine_intro,
    "lh": engine_lh,
}

BANK_COLUMNS = [
    "Topic", "Text", "a", "b", "c", "d", "extras", "correct",
    "idea_id", "block_id", "q_id", "b_info",
]


def read_xlsx(path):
    """Read the first sheet of an xlsx test bank into a list of row dicts,
    mirroring readxl::read_xlsx (header row = column names, empty cells = NA).
    Cell values are converted with r_char so later paste0/comparison behavior
    matches R's as.character semantics."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    header = [h for h in header if h is not None]
    out = []
    for raw in rows:
        # readxl drops rows that are entirely empty
        if all(v is None for v in raw[: len(header)]):
            continue
        row = {}
        for name, value in zip(header, raw):
            value = r_char(value)
            if isinstance(value, str):
                # readxl defaults: trim_ws = TRUE, na = "" (whitespace-only
                # or empty cells become NA)
                value = value.strip(" \t\r\n")
                if value == "":
                    value = None
            row[name] = value
        # Guarantee every expected column exists (missing -> NA)
        for col in BANK_COLUMNS:
            row.setdefault(col, None)
        out.append(row)
    wb.close()
    return out


def load_full_bank(banks_to_load, base_dir="."):
    """Port of the purrr::map_df loop in app.R: read each bank and overwrite
    its Topic column with the configured topic label.

    banks_to_load: list of {"path": ..., "topic": ...} dicts (enabled banks
    only, in the order they appear in the config)."""
    full_bank = []
    for bank in banks_to_load:
        path = os.path.join(base_dir, bank["path"])
        rows = read_xlsx(path)
        for row in rows:
            row["Topic"] = bank["topic"]
        full_bank.extend(rows)
    return full_bank


def unique_topics(full_bank):
    """unique(full_bank$Topic), preserving first-appearance order."""
    seen = []
    for row in full_bank:
        if row["Topic"] not in seen:
            seen.append(row["Topic"])
    return seen


def generate_exam(full_bank, sel_topics, n_questions, seed_input, landscape,
                  engine, semester, classnum, exam_dir="./Exam",
                  real_exam=False, real_exam_title=None,
                  num_versions_input=1, title_input="Practice Exam:"):
    """Port of the observeEvent(input$generate_exam, ...) handler.

    Returns {"status": "ok", "tex": ...} on success or
    {"status": "error", "message": ...} mirroring the Shiny notifications.
    """
    # req(input$sel_topics): with no topics selected, Shiny's req() halts the
    # handler silently and nothing happens.
    if not sel_topics:
        return {"status": "noop"}

    if not os.path.isdir(exam_dir):
        os.makedirs(exam_dir)

    # Use random seed if provided (set.seed(as.numeric(input$seed_input)))
    if seed_input != "":
        try:
            seed = float(seed_input)
        except ValueError:
            # R: as.numeric("abc") is NA and set.seed(NA) errors out
            return {
                "status": "error",
                "message": "Random seed must be a number.",
            }
        exam_common.set_seed(seed)
    else:
        exam_common.reseed_from_system()

    # Filter and sample
    cur_bank = [row for row in full_bank if row["Topic"] in sel_topics]
    if len(cur_bank) < n_questions:
        return {
            "status": "error",
            "message": (
                "Not enough questions in selected topics. Maximum available: "
                + str(len(cur_bank)) + " questions."
            ),
        }

    qind = exam_common.r_sample_permutation(len(cur_bank))
    # Instantiate parameterized math questions (banks with a `params`
    # column / [[expr]] placeholders): every generation draws fresh
    # numbers from the same seeded RNG, so a seed reproduces the exam.
    # Static banks pass through byte-identically with no RNG draws.
    try:
        cur_bank = param_engine.instantiate_bank(cur_bank, exam_common.get_rng())
    except param_engine.ParamError as e:
        return {"status": "error", "message": str(e)}
    # The examdesign class renders num_versions versions in one document
    # (rearranged question order per version, front page shows the version
    # letter) — the same mechanism app.R used for real exams, so professors
    # can generate actual exams with identical formatting.
    num_versions = num_versions_input
    title = title_input

    if real_exam:
        num_versions = 3
        title = real_exam_title

    tex_path = os.path.join(exam_dir, "Practice_Exam.tex")
    tex = ENGINES[engine].create_exam(
        bank=cur_bank,
        qind=qind,
        tit=title,
        sem=semester,
        classnum=classnum,
        numVer=num_versions,
        file=tex_path,
        frontPage=True,
        max_q=n_questions,
        real_exam=real_exam,
        landscape=landscape,
    )

    return {"status": "ok", "tex": tex}
